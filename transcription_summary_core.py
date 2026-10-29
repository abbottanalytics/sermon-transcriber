#!/usr/bin/env python3

# -------------------------------------------------
# PURPOSE:
# Shared core utilities for transcription processing.
#
# This module provides:
# - LLM interaction
# - transcript chunking
# - docx/text writing
# - transcript cleaning
#
# DESIGN PRINCIPLES:
# - No meeting/sermon-specific logic
# - Reusable across multiple workflows
# - Safe for large transcripts via chunking
# -------------------------------------------------

import re
import time
from pathlib import Path
from typing import List

try:
    import tiktoken
except ImportError:
    tiktoken = None

from openai import OpenAI
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

MODEL = "gpt-4.1-mini"
# Transcript cleaning goes in chunks of about this many tokens. Given much
# more at once, the model condenses (a 7,000-word class came back as 2,100
# words); at this size it tidies instead.
CLEAN_CHUNK_TOKENS = 3000
SLEEP_BETWEEN_CALLS = 0.4


# -------------------------------------------------
# METADATA HELPERS
# -------------------------------------------------

def build_metadata_header(fields: dict) -> str:
    """
    Build a markdown metadata block to prepend to document content.

    Example:
        build_metadata_header({"Speaker": "Pastor Joel Wood", "Event": "Morning Worship"})
        → "**Speaker:** Pastor Joel Wood\n**Event:** Morning Worship\n\n---\n\n"
    """
    lines = [f"**{k}:** {v}" for k, v in fields.items() if v]
    return ("\n".join(lines) + "\n\n---\n\n") if lines else ""


# -------------------------------------------------
# FILE OUTPUT HELPERS
# -------------------------------------------------

# **bold**, __bold__, *italic* or _italic_ inside a line
_INLINE = re.compile(r"(\*\*.+?\*\*|__.+?__|(?<![*\w])\*(?!\s).+?(?<!\s)\*(?![*\w])|(?<!\w)_(?!\s).+?(?<!\s)_(?!\w))")
_BULLET = re.compile(r"^(\s*)[-*•]\s+(.*)$")


def _add_inline(paragraph, text: str):
    """Add text to a paragraph, turning Markdown bold/italic marks into formatting."""
    for part in _INLINE.split(text):
        if not part:
            continue
        if (part.startswith("**") and part.endswith("**") or
                part.startswith("__") and part.endswith("__")) and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part[0] in "*_" and part[-1] == part[0] and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _add_rule(doc):
    """A horizontal line, for a Markdown '---'."""
    p = doc.add_paragraph()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "999999")):
        bottom.set(qn(key), value)
    border.append(bottom)
    p._p.get_or_add_pPr().append(border)


def write_docx(path, markdown_text, overwrite):
    if path.exists() and not overwrite:
        print(f"SKIP (exists): {path}")
        return

    doc = Document()

    for line in markdown_text.splitlines():
        line = line.rstrip()            # Markdown's trailing-space line breaks
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        bullet = _BULLET.match(line)
        if heading:
            # Headings are bold already; drop any marks inside them
            doc.add_heading(re.sub(r"\*\*|__", "", heading.group(2)), level=len(heading.group(1)))
        elif re.fullmatch(r"\s*([-*_])(\s*\1){2,}\s*", line):
            _add_rule(doc)
        elif line.startswith("> "):
            p = doc.add_paragraph()
            _add_inline(p, line[2:])
            p.paragraph_format.left_indent = Pt(18)
        elif bullet:
            # Indented bullets (sub-points) go one level in
            level = 2 if len(bullet.group(1).expandtabs(4)) >= 2 else 1
            p = doc.add_paragraph(style="List Bullet" if level == 1 else "List Bullet 2")
            _add_inline(p, bullet.group(2))
        elif line.strip() == "":
            doc.add_paragraph("")
        else:
            p = doc.add_paragraph()
            _add_inline(p, line)

    doc.save(path)
    print(f"WROTE: {path}")


def write_text(path: Path, content: str, overwrite: bool):
    if path.exists() and not overwrite:
        print(f"SKIP (exists): {path}")
        return
    path.write_text(content.strip() + "\n", encoding="utf-8")
    print(f"WROTE: {path}")


# -------------------------------------------------
# TOKEN + CHUNKING
# -------------------------------------------------

def estimate_tokens(text: str) -> int:
    if not tiktoken:
        return len(text) // 4
    enc = tiktoken.encoding_for_model(MODEL)
    return len(enc.encode(text))


def chunk_text(text: str, max_tokens: int) -> List[str]:
    """
    Split text into chunks of about max_tokens, ending at sentence breaks.

    WHY:
    - Each chunk is cleaned separately, and the cleaned chunks are joined, so
      chunks don't overlap (overlap would repeat sentences in the output)
    - Breaking between sentences keeps each chunk's text whole
    """
    max_chars = max_tokens * 4
    sentences = re.split(r"(?<=[.!?])\s+", " ".join(text.split()))
    chunks, current = [], ""
    for sentence in sentences:
        if current and len(current) + len(sentence) + 1 > max_chars:
            chunks.append(current)
            current = ""
        current = f"{current} {sentence}" if current else sentence
    if current:
        chunks.append(current)
    return chunks


# -------------------------------------------------
# INPUT HELPERS
# -------------------------------------------------

def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore").strip()


# -------------------------------------------------
# LLM CALL (WITH RETRY + SAFE OUTPUT)
# -------------------------------------------------

def call_llm(client: OpenAI, system: str, user: str, retries=3) -> str:
    """
    Safe LLM call wrapper.

    Adds:
    - retry logic
    - robust output extraction
    """

    for attempt in range(retries):
        try:
            resp = client.responses.create(
                model=MODEL,
                input=[
                    {"role": "system", "content": system},
                    {"role": "user", "content": user},
                ],
            )

            # SAFER OUTPUT EXTRACTION
            if hasattr(resp, "output_text") and resp.output_text:
                return resp.output_text

            # fallback (rare but safer)
            return resp.output[0].content[0].text

        except Exception as e:
            print(f"LLM error (attempt {attempt+1}): {e}")
            time.sleep(1 + attempt)

    raise RuntimeError("LLM call failed after retries")


# -------------------------------------------------
# TRANSCRIPT CLEANING
# -------------------------------------------------

def clean_transcript(client: OpenAI, raw_text: str) -> str:
    """
    Clean raw ASR transcript into readable structured text.

    DESIGN CHOICES:
    - No summarization
    - Preserve meaning
    - Add light structure (headings)
    - Avoid speaker hallucination
    """

    system = (
        "You are editing a transcript produced by automatic speech recognition.\n"
        "- Preserve meaning; do not summarize, shorten or paraphrase. Keep every point, "
        "example and story, in the speaker's own words.\n"
        "- Fix punctuation and structure.\n"
        "- Remove obvious filler words.\n"
        "- Add section headings (##) when topics shift.\n"
        "- Do NOT guess speaker identities.\n"
    )

    chunks = chunk_text(raw_text, CLEAN_CHUNK_TOKENS)
    parts = []

    for i, chunk in enumerate(chunks, 1):
        print(f"  Cleaning chunk {i}/{len(chunks)}")
        parts.append(call_llm(client, system, chunk))
        time.sleep(SLEEP_BETWEEN_CALLS)

    return "\n\n".join(parts).strip()

